"""nxdoc - build Word documents in the Nutanix Services document style.

Fully self-contained: no template file needed. Styles, list numbering and the
header line are embedded below; Montserrat is fetched from npm and embedded in
the .docx so it looks right on any PC. Em dashes are never written.
"""
import datetime as _dt
import glob
import io
import os
import re
import shutil
import subprocess
import tarfile
import tempfile
import uuid
import base64

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches

# Page geometry per variant (twentieths of a point).
VARIANTS = {
    "guide": dict(top=1152, bottom=1152, left=720, right=720, small_tables=False),
    "asbuilt": dict(top=1440, bottom=1440, left=1440, right=1440, small_tables=True),
}
PAGE_W, PAGE_H = 12240, 15840

BULLET_STYLES = ["List Bullet", "List Bullet 2", "List Bullet 3", "List Bullet 4"]
NUMBER_STYLES = ["List Number", "List Number 2", "List Number 3"]

# Montserrat from npm (Google Fonts TTFs); embedded into every document.
FONT_PKG = "@expo-google-fonts/montserrat@0.4.2"
FONT_FILES = {  # family -> {embed slot: file stem}
    "Montserrat": {"embedRegular": "400Regular", "embedBold": "700Bold",
                   "embedItalic": "400Regular_Italic", "embedBoldItalic": "700Bold_Italic"},
    "Montserrat SemiBold": {"embedRegular": "600SemiBold", "embedItalic": "600SemiBold_Italic"},
    "Montserrat Medium": {"embedRegular": "500Medium"},
}
FONT_CACHE = os.path.expanduser("~/.cache/nxdoc/fonts")

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def _w(tag):
    return qn("w:" + tag)


def _x(xml, *prefixes):
    """parse an XML snippet, adding namespace declarations to its root tag"""
    return parse_xml(re.sub(r"^<([\w:]+)", r"<\1 " + nsdecls(*(prefixes or ("w",))), xml, count=1))


class NxDoc:
    def __init__(self, variant="guide", title="Document Title",
                 subtitle="Infrastructure Deployment or Expansion", author=None, margins=None):
        if variant not in VARIANTS:
            raise ValueError("variant must be one of %s" % list(VARIANTS))
        self.cfg = dict(VARIANTS[variant])
        if margins:
            self.cfg.update(margins)
        self.variant = variant
        self.title, self.subtitle, self.author = _no_dash(title), _no_dash(subtitle or ""), author
        self.text_width = PAGE_W - self.cfg["left"] - self.cfg["right"]
        self.doc = Document()
        self.body = self.doc.element.body
        self._headings = []
        self._fig = self._tbl = 0
        self._bm = 9000
        self._install_styles()
        self._page_setup()
        self._build_cover()
        self._build_toc_block()
        self._build_header_footer()

    # ------------------------------------------------------------------ base
    def _install_styles(self):
        st = self.doc.styles.element
        ids = {re.search(r'w:styleId="([^"]+)"', x).group(1) for x in STYLES}
        for s in list(st.findall(_w("style"))):
            if s.get(_w("styleId")) in ids:
                st.remove(s)
        old = st.find(_w("docDefaults"))
        new = _x(DOC_DEFAULTS)
        if old is not None:
            old.addprevious(new); st.remove(old)
        else:
            st.insert(0, new)
        for x in STYLES:
            st.append(_x(x))
        num = self.doc.part.numbering_part.element
        for c in list(num):
            num.remove(c)
        for x in ABSTRACT_NUMS:
            num.append(_x(x))
        for x in NUMS:
            num.append(_x(x))

    def _page_setup(self):
        for k in list(self.body):
            if k.tag != _w("sectPr"):
                self.body.remove(k)
        sp = self.body.find(_w("sectPr"))
        sp.find(_w("pgSz")).set(_w("w"), str(PAGE_W)); sp.find(_w("pgSz")).set(_w("h"), str(PAGE_H))
        pm = sp.find(_w("pgMar"))
        for k in ("top", "bottom", "left", "right"):
            pm.set(_w(k), str(self.cfg[k]))
        pm.set(_w("header"), "432"); pm.set(_w("footer"), "288")
        self.doc.sections[0].different_first_page_header_footer = True
        # same Word compatibility settings as the original documents (Word 2013+ layout)
        settings = self.doc.settings.element
        z = settings.find(_w("zoom"))
        if z is not None:
            z.set(_w("percent"), "100")
        old = settings.find(_w("compat"))
        new = _x('<w:compat>' + "".join(
            '<w:compatSetting w:name="%s" w:uri="http://schemas.microsoft.com/office/word" w:val="%s"/>' % kv
            for kv in (("compatibilityMode", "15"), ("overrideTableStyleFontSizeAndJustification", "1"),
                       ("enableOpenTypeFeatures", "1"), ("doNotFlipMirrorIndents", "1"),
                       ("differentiateMultirowTableHeaders", "1"))) + '</w:compat>')
        if old is not None:
            old.addprevious(new); settings.remove(old)

    def _build_cover(self):
        """All-black text cover: NUTANIX (Montserrat Bold 72pt), title, subtitle, date."""
        today = _dt.date.today().strftime("%B %Y")
        rp = '<w:rPr><w:color w:val="000000"/><w:sz w:val="24"/></w:rPr>'
        paras = [
            '<w:p><w:pPr><w:spacing w:before="3600" w:after="480"/><w:ind w:left="0"/></w:pPr>'
            '<w:r><w:rPr><w:rFonts w:ascii="Montserrat" w:hAnsi="Montserrat" w:cs="Montserrat"/><w:b/><w:bCs/>'
            '<w:color w:val="000000"/><w:sz w:val="144"/><w:szCs w:val="144"/></w:rPr><w:t>NUTANIX</w:t></w:r></w:p>',
            '<w:p><w:pPr><w:pStyle w:val="Title"/><w:spacing w:before="0" w:after="0"/></w:pPr>'
            '<w:r><w:rPr><w:rFonts w:ascii="Montserrat SemiBold" w:hAnsi="Montserrat SemiBold"/><w:b w:val="0"/>'
            '<w:color w:val="000000"/><w:sz w:val="56"/><w:szCs w:val="56"/></w:rPr>'
            '<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % _esc(self.title),
        ]
        if self.subtitle:
            paras.append('<w:p><w:pPr><w:pStyle w:val="Subtitle"/><w:spacing w:before="120" w:after="0"/></w:pPr>'
                         '<w:r><w:rPr><w:color w:val="000000"/><w:sz w:val="32"/><w:szCs w:val="32"/></w:rPr>'
                         '<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % _esc(self.subtitle))
        paras.append(
            '<w:p><w:pPr><w:spacing w:before="480" w:after="0"/><w:ind w:left="0"/></w:pPr>'
            '<w:r>%s<w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r>%s<w:instrText xml:space="preserve"> DATE \\@ "MMMM yyyy" </w:instrText></w:r>'
            '<w:r>%s<w:fldChar w:fldCharType="separate"/></w:r><w:r>%s<w:t>%s</w:t></w:r>'
            '<w:r>%s<w:fldChar w:fldCharType="end"/></w:r><w:r><w:br w:type="page"/></w:r></w:p>'
            % (rp, rp, rp, rp, today, rp))
        sdt = _x('<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Cover Pages"/><w:docPartUnique/>'
                 '</w:docPartObj></w:sdtPr><w:sdtContent>%s</w:sdtContent></w:sdt>'
                 % "".join(paras))
        self.body.find(_w("sectPr")).addprevious(sdt)

    def _build_toc_block(self):
        sdt = _x('<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/>'
                 '</w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr>'
                 '<w:r><w:t>Contents</w:t></w:r></w:p></w:sdtContent></w:sdt>')
        self.body.find(_w("sectPr")).addprevious(sdt)
        self._toc_content = sdt.find(_w("sdtContent"))
        self._toc_heading = self._toc_content[0]

    def _build_header_footer(self):
        s = self.doc.sections[0]
        # running header: purple chevron line + "<subtitle>: <title>" as plain text
        hdr = s.header
        rid, _ = hdr.part.get_or_add_image(io.BytesIO(base64.b64decode("".join(HEADER_LINE_PNG))))
        el = hdr._element
        for c in list(el):
            el.remove(c)
        text = "%s: %s" % (self.subtitle, self.title) if self.subtitle else self.title
        el.append(_x(
            '<w:p><w:pPr><w:pStyle w:val="Header"/></w:pPr><w:r><w:rPr><w:noProof/></w:rPr><w:drawing>'
            '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658242" behindDoc="1" '
            'locked="1" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="page"><wp:posOffset>27622</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="page"><wp:posOffset>447675</wp:posOffset></wp:positionV>'
            '<wp:extent cx="7717155" cy="246380"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            '<wp:docPr id="1001" name="Header line"/><wp:cNvGraphicFramePr/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            '<pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="line.png"/><pic:cNvPicPr/></pic:nvPicPr>'
            '<pic:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="7717155" cy="246380"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>'
            '</wp:anchor></w:drawing></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t xml:space="preserve">%s</w:t>'
            '</w:r></w:p>' % (rid, _esc(text)), "w", "wp", "a", "pic", "r"))
        # footer: centered page number only; cover page has no header/footer
        ft = s.footer._element
        for c in list(ft):
            ft.remove(c)
        rp = '<w:rPr><w:sz w:val="18"/></w:rPr>'
        ft.append(_x('<w:p><w:pPr><w:pStyle w:val="Footer"/><w:tabs><w:tab w:val="clear" w:pos="10800"/></w:tabs>'
                     '<w:ind w:left="0"/><w:jc w:val="center"/></w:pPr><w:r>%s<w:fldChar w:fldCharType="begin"/></w:r>'
                     '<w:r>%s<w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
                     '<w:r>%s<w:fldChar w:fldCharType="separate"/></w:r><w:r>%s<w:t>2</w:t></w:r>'
                     '<w:r>%s<w:fldChar w:fldCharType="end"/></w:r></w:p>' % (rp, rp, rp, rp, rp)))
        for hf in (s.first_page_header, s.first_page_footer):
            e = hf._element
            for c in list(e):
                e.remove(c)
            e.append(_x('<w:p><w:pPr><w:pStyle w:val="Header"/></w:pPr></w:p>'))

    # ---------------------------------------------------------------- inline
    def _add_inline(self, par, text, bold=False):
        """Add text with **bold**, *italic*, `code` and [link](url) markup."""
        for part in _INLINE.split(_no_dash(text)):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                r = par.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1]); r.style = self.doc.styles["Code: Inline"]
                if bold: r.bold = True
            elif part.startswith("[") and "](" in part and part.endswith(")"):
                label, url = part[1:].split("](", 1)
                self._hyperlink(par, label, url[:-1])
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = par.add_run(part[1:-1]); r.italic = True
                if bold: r.bold = True
            else:
                r = par.add_run(part)
                if bold: r.bold = True
        return par

    def _hyperlink(self, par, label, url):
        rid = self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
        par._p.append(_x('<w:hyperlink r:id="%s" w:history="1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr>'
                         '<w:t xml:space="preserve">%s</w:t></w:r></w:hyperlink>' % (rid, _esc(label)), "w", "r"))

    # ------------------------------------------------------------- content
    def _heading(self, level, text):
        text = _no_dash(text)
        p = self.doc.add_paragraph(style="Heading %d" % level)
        p.add_run(text)
        self._headings.append((level, text, p._p))
        return p

    def h1(self, text):
        """Numbered chapter heading (starts a new page)."""
        return self._heading(1, text)

    def h2(self, text):
        return self._heading(2, text)

    def h3(self, text):
        return self._heading(3, text)

    def p(self, text="", bold=False, style="Normal"):
        return self._add_inline(self.doc.add_paragraph(style=style), text, bold=bold)

    def note(self, text):
        """Grey shaded callout (style 'Note')."""
        return self.p(text, style="Note")

    def instructions(self, text):
        """Yellow 'CONSULTANT - ...' guidance line (style 'Instructions')."""
        return self.p(text, style="Instructions")

    def code(self, lines):
        """Boxed Courier block, kept verbatim. `lines`: string or list."""
        if isinstance(lines, str):
            lines = lines.split("\n")
        for ln in lines:
            self.doc.add_paragraph(style="Code: Block").add_run(ln)

    def page_break(self):
        from docx.enum.text import WD_BREAK
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # lists ----------------------------------------------------------------
    def _restart_num(self, style_name):
        """New w:num on the style's list definition, restarting at 1."""
        numpr = self.doc.styles[style_name].element.find("./" + _w("pPr") + "/" + _w("numPr"))
        nid = numpr.find(_w("numId")).get(_w("val"))
        il = numpr.find(_w("ilvl"))
        ilvl = int(il.get(_w("val"))) if il is not None else 0
        numbering = self.doc.part.numbering_part.element
        abs_id = next(int(n.find(_w("abstractNumId")).get(_w("val"))) for n in numbering.findall(_w("num"))
                      if n.get(_w("numId")) == nid)
        new = numbering.add_num(abs_id)
        new.add_lvlOverride(ilvl).add_startOverride(1)
        return new.numId, ilvl

    def _list(self, items, kind, level):
        styles = BULLET_STYLES if kind == "bullet" else NUMBER_STYLES
        style = styles[min(level, len(styles) - 1)]
        num = self._restart_num(style) if kind == "number" else None
        for it in items:
            sub, sub_kind = None, kind
            if isinstance(it, (tuple, list)):
                text = it[0]
                sub = it[1] if len(it) > 1 else None
                if len(it) > 2:
                    sub_kind = it[2]
            else:
                text = it
            par = self.doc.add_paragraph(style=style)
            if num:
                numPr = par._p.get_or_add_pPr().get_or_add_numPr()
                numPr.get_or_add_ilvl().val = num[1]
                numPr.get_or_add_numId().val = num[0]
            self._add_inline(par, text)
            if sub:
                self._list(sub, sub_kind, level + 1)

    def bullets(self, items, level=0):
        """items: str | (text, [children]) | (text, [children], 'number'|'bullet')"""
        self._list(items, "bullet", level)

    def numbered(self, items, level=0):
        """Numbered steps; numbering restarts at 1 for every call."""
        self._list(items, "number", level)

    # tables ---------------------------------------------------------------
    def table(self, headers, rows, widths=None, small=None, caption=None, indent=360, bold_first_col=False):
        """Purple-header table. widths relative (optional). Caption goes ABOVE."""
        small = self.cfg["small_tables"] if small is None else small
        if caption:
            self._caption("Table", caption)
        ncols = len(headers)
        total = self.text_width - indent
        if widths:
            s = float(sum(widths))
            cw = [int(total * w / s) for w in widths]
        else:
            cw = [total // ncols] * ncols
        cw[-1] = total - sum(cw[:-1])
        tbl = self.doc.add_table(rows=1 + len(rows), cols=ncols)
        tbl.style = self.doc.styles["Nutanix Table: Small" if small else "Nutanix Table"]
        cell_style = "Table: Text Small" if small else "Table: Text"
        tblPr = tbl._tbl.tblPr
        for tag in ("tblW", "tblInd"):
            e = tblPr.find(_w(tag))
            if e is not None:
                tblPr.remove(e)
        tw = OxmlElement("w:tblW"); tw.set(_w("w"), str(total)); tw.set(_w("type"), "dxa")
        ti = OxmlElement("w:tblInd"); ti.set(_w("w"), str(indent)); ti.set(_w("type"), "dxa")
        tblPr.find(_w("tblStyle")).addnext(tw); tw.addnext(ti)
        lay = OxmlElement("w:tblLayout"); lay.set(_w("type"), "fixed"); ti.addnext(lay)
        for gc, w in zip(tbl._tbl.tblGrid.findall(_w("gridCol")), cw):
            gc.set(_w("w"), str(w))
        tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for ri, row in enumerate([headers] + [list(r) for r in rows]):
            for ci in range(ncols):
                cell = tbl.cell(ri, ci)
                cell.width = Emu(cw[ci] * 635)
                val = row[ci] if ci < len(row) else ""
                par = cell.paragraphs[0]
                par.style = self.doc.styles[cell_style]
                for li, ln in enumerate(str(val if val is not None else "").split("\n")):
                    if li:
                        par = cell.add_paragraph(style=cell_style)
                    self._add_inline(par, ln, bold=(bold_first_col and ci == 0 and ri > 0))
        self.doc.add_paragraph()  # spacer after table, as in the source docs
        return tbl

    # figures --------------------------------------------------------------
    def _caption(self, kind, text):
        if kind == "Figure":
            self._fig += 1; n = self._fig
        else:
            self._tbl += 1; n = self._tbl
        par = self.doc.add_paragraph(style="Caption")
        par.add_run(kind + " ")
        for x in ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                  '<w:r><w:instrText xml:space="preserve"> SEQ %s \\* ARABIC </w:instrText></w:r>' % kind,
                  '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', '<w:r><w:t>%d</w:t></w:r>' % n,
                  '<w:r><w:fldChar w:fldCharType="end"/></w:r>'):
            par._p.append(_x(x))
        par.add_run(" " + _no_dash(text))
        return par

    def figure(self, image_path, caption=None, width_in=None):
        """Centered image (style 'Figure'); caption goes BELOW ('Figure N ...')."""
        max_w = (self.text_width - 360) / 1440.0
        par = self.doc.add_paragraph(style="Figure")
        if width_in is None:
            try:
                from PIL import Image
                with Image.open(image_path) as im:
                    dpi = im.info.get("dpi", (96, 96))[0] or 96
                    width_in = min(max_w, im.width / float(dpi))
            except Exception:
                width_in = max_w
        par.add_run().add_picture(image_path, width=Inches(min(width_in, max_w)))
        if caption:
            self._caption("Figure", caption)
        return par

    def placeholder_figure(self, caption="<Insert Caption here>"):
        """'<Insert Figure here>' placeholder used in templates/as-builts."""
        self.doc.add_paragraph("<Insert Figure here>", style="Figure")
        return self.doc.add_paragraph(caption, style="Caption")

    # ----------------------------------------------------------------- TOC
    def _numbers(self):
        c, out = [0, 0, 0], []
        for level, _, _ in self._headings:
            c[level - 1] += 1
            for i in range(level, 3):
                c[i] = 0
            out.append(".".join(str(x) for x in c[:level]))
        return out

    def _bookmark_headings(self):
        names = []
        for _, _, el in self._headings:
            self._bm += 1
            name = "_Toc%09d" % (190000000 + self._bm)
            bs = OxmlElement("w:bookmarkStart"); bs.set(_w("id"), str(self._bm)); bs.set(_w("name"), name)
            be = OxmlElement("w:bookmarkEnd"); be.set(_w("id"), str(self._bm))
            ppr = el.find(_w("pPr"))
            if ppr is not None:
                ppr.addnext(bs)
            else:
                el.insert(0, bs)
            el.append(be)
            names.append(name)
        return names

    def _build_toc(self, names, pages):
        content = self._toc_content
        for p in list(content):
            if p is not self._toc_heading:
                content.remove(p)
        num_tab = {1: 360, 2: 960, 3: 1440}
        paras = []
        for (lv, tx, _), nm, nu, pg in zip(self._headings, names, self._numbers(), pages):
            p = _x('<w:p><w:pPr><w:pStyle w:val="TOC%d"/><w:tabs><w:tab w:val="left" w:pos="%d"/>'
                   '<w:tab w:val="clear" w:pos="11995"/><w:tab w:val="right" w:leader="dot" w:pos="%d"/></w:tabs>'
                   '<w:spacing w:before="%d" w:after="60"/></w:pPr></w:p>'
                   % (lv, num_tab[lv], self.text_width, 160 if lv == 1 else 60))
            if not paras:
                for x in ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                          '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>',
                          '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'):
                    p.append(_x(x))
            rp, rh = '<w:rPr><w:noProof/></w:rPr>', '<w:rPr><w:noProof/><w:webHidden/></w:rPr>'
            p.append(_x(
                '<w:hyperlink w:anchor="%s" w:history="1"><w:r>%s<w:t>%s</w:t></w:r><w:r>%s<w:tab/></w:r>'
                '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r><w:r>%s<w:tab/></w:r>'
                '<w:r>%s<w:fldChar w:fldCharType="begin"/></w:r>'
                '<w:r>%s<w:instrText xml:space="preserve"> PAGEREF %s \\h </w:instrText></w:r>'
                '<w:r>%s<w:fldChar w:fldCharType="separate"/></w:r><w:r>%s<w:t>%s</w:t></w:r>'
                '<w:r>%s<w:fldChar w:fldCharType="end"/></w:r></w:hyperlink>'
                % (nm, rp, _esc(nu), rp, rp, _esc(tx), rh, rh, rh, nm, rh, rh, pg, rh)))
            paras.append(p)
        if paras:
            paras[-1].append(_x('<w:r><w:fldChar w:fldCharType="end"/></w:r>'))
        for p in paras:
            content.append(p)

    # ----------------------------------------------------------------- save
    def save(self, path, update_fields_on_open=True, page_numbers=True, embed_fonts=True):
        """Write the .docx: TOC with estimated page numbers, embedded Montserrat.
        update_fields_on_open: Word refreshes the TOC on open (one-time prompt)."""
        names = self._bookmark_headings()
        self._build_toc(names, ["" for _ in names])
        self._finish_props(update_fields_on_open)
        for rId, rel in list(self.doc.part.rels.items()):
            if "customXml" in rel.reltype:
                del self.doc.part.rels[rId]
        if embed_fonts:
            try:
                self._embed_fonts()
            except Exception as e:
                print("nxdoc: fonts not embedded (%s); Word will use installed Montserrat" % e)
        self.doc.save(path)
        if page_numbers and names:
            try:
                self._build_toc(names, _estimate_pages(path, [t for _, t, _ in self._headings]))
                self.doc.save(path)
            except Exception as e:
                print("nxdoc: page-number estimation skipped (%s)" % e)
        return path

    def _embed_fonts(self):
        fdir = fetch_fonts()
        ft_part = next(r.target_part for r in self.doc.part.rels.values() if r.reltype == RT.FONT_TABLE)
        ft = parse_xml(ft_part.blob)
        n = 0
        for fam, slots in FONT_FILES.items():
            for f in ft.findall(_w("font")):
                if f.get(_w("name")) == fam:
                    ft.remove(f)
            fe = _x('<w:font w:name="%s"><w:charset w:val="00"/><w:family w:val="auto"/>'
                    '<w:pitch w:val="variable"/></w:font>' % fam)
            for slot, stem in slots.items():
                data = bytearray(open(os.path.join(fdir, "Montserrat_%s.ttf" % stem), "rb").read())
                guid = str(uuid.uuid4()).upper()
                key = bytes.fromhex(guid.replace("-", ""))[::-1]
                for i in range(32):
                    data[i] ^= key[i % 16]
                n += 1
                part = Part(PackURI("/word/fonts/nxfont%d.odttf" % n),
                            "application/vnd.openxmlformats-officedocument.obfuscatedFont",
                            bytes(data), self.doc.part.package)
                rid = ft_part.relate_to(part, RT.FONT)
                fe.append(_x('<w:%s r:id="%s" w:fontKey="{%s}"/>' % (slot, rid, guid), "w", "r"))
            ft.append(fe)
        from lxml import etree
        ft_part._blob = etree.tostring(ft, xml_declaration=True, encoding="UTF-8", standalone=True)
        settings = self.doc.settings.element
        if settings.find(_w("embedTrueTypeFonts")) is None:
            e = OxmlElement("w:embedTrueTypeFonts")
            first = next((c for c in settings if c.tag not in {_w(t) for t in (
                "writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime",
                "doNotDisplayPageBoundaries", "displayBackgroundShape", "printPostScriptOverText",
                "printFractionalCharacterWidth", "printFormsData")}), None)
            first.addprevious(e) if first is not None else settings.append(e)

    def _finish_props(self, update_fields):
        cp = self.doc.core_properties
        cp.title, cp.subject = self.title, self.subtitle
        cp.author = cp.last_modified_by = self.author or ""
        cp.revision = 1
        cp.created = cp.modified = _dt.datetime.now()
        settings = self.doc.settings.element
        uf = settings.find(_w("updateFields"))
        if update_fields and uf is None:
            uf = OxmlElement("w:updateFields"); uf.set(_w("val"), "true")
            after = {_w(a) for a in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                                     "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                                     "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                                     "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                                     "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")}
            nxt = next((c for c in settings if c.tag in after), None)
            nxt.addprevious(uf) if nxt is not None else settings.append(uf)
        elif not update_fields and uf is not None:
            settings.remove(uf)


# ======================================================================= utils
def _no_dash(s):
    """House rule: no em dashes. ' — ' -> ', ' and a bare — -> ' - '."""
    if not s:
        return s
    return re.sub(r"\s*—\s*", lambda m: ", " if m.group(0).strip() != m.group(0) else " - ", str(s))


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def fetch_fonts():
    """Download Montserrat TTFs once (npm registry) into FONT_CACHE and install
    them for LibreOffice previews. Returns the cache directory."""
    need = {stem for slots in FONT_FILES.values() for stem in slots.values()}
    have = lambda: all(os.path.exists(os.path.join(FONT_CACHE, "Montserrat_%s.ttf" % s)) for s in need)
    if not have():
        os.makedirs(FONT_CACHE, exist_ok=True)
        tmp = tempfile.mkdtemp(prefix="nxfont_")
        subprocess.run(["npm", "pack", FONT_PKG, "--silent"], cwd=tmp, capture_output=True, timeout=180)
        tgz = glob.glob(os.path.join(tmp, "*.tgz"))
        if not tgz:
            raise RuntimeError("could not download %s" % FONT_PKG)
        with tarfile.open(tgz[0]) as t:
            for m in t.getmembers():
                name = os.path.basename(m.name)
                if name.endswith(".ttf") and name[len("Montserrat_"):-4] in need:
                    with open(os.path.join(FONT_CACHE, name), "wb") as f:
                        f.write(t.extractfile(m).read())
        shutil.rmtree(tmp, ignore_errors=True)
        if not have():
            raise RuntimeError("font files missing")
    fdir = os.path.expanduser("~/.fonts/nxdoc")
    if not os.path.isdir(fdir):
        shutil.copytree(FONT_CACHE, fdir)
        subprocess.run(["fc-cache", "-f", fdir], capture_output=True)
    return FONT_CACHE


def _soffice():
    cand = glob.glob("/mnt/skills/public/docx/scripts/office/soffice.py")
    if cand:
        return ["python3", cand[0]]
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice not available")
    return [exe]


def render(docx_path, outdir=None, dpi=60, pages=None):
    """Convert to PDF (+ JPG pages) with LibreOffice. Returns (pdf, [jpgs])."""
    try:
        fetch_fonts()
    except Exception:
        pass
    outdir = outdir or tempfile.mkdtemp(prefix="nxdoc_")
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(_soffice() + ["--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   capture_output=True, timeout=300)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if not os.path.exists(pdf):
        raise RuntimeError("PDF conversion failed")
    jpgs = []
    if dpi:
        base = os.path.join(outdir, "page")
        cmd = ["pdftoppm", "-jpeg", "-r", str(dpi)]
        if pages:
            cmd += ["-f", str(pages[0]), "-l", str(pages[1])]
        subprocess.run(cmd + [pdf, base], capture_output=True)
        jpgs = sorted(glob.glob(base + "-*.jpg"))
    return pdf, jpgs


def _estimate_pages(docx_path, heading_texts):
    tmp = tempfile.mkdtemp(prefix="nxdoc_pg_")
    pdf, _ = render(docx_path, tmp, dpi=0)
    pages = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout.split("\f")
    norm = lambda s: re.sub(r"[^0-9a-z]", "", s.lower())
    npages = [norm(p) for p in pages]
    start = 1  # skip cover + contents pages (dot leaders)
    while start < len(pages) and re.search(r"\.{6,}", pages[start]):
        start += 1
    res, cur = [], start
    for h in heading_texts:
        key = norm(h)
        found = next((i for i in range(cur, len(npages)) if key and key in npages[i]), None)
        if found is None:
            res.append("")
        else:
            res.append(str(found + 1)); cur = found
    shutil.rmtree(tmp, ignore_errors=True)
    return res


def make_contact_sheet(jpgs, out_path, cols=4):
    from PIL import Image
    ims = [Image.open(j) for j in jpgs]
    if not ims:
        return None
    w, h = ims[0].size
    sheet = Image.new("RGB", (w * cols, h * ((len(ims) + cols - 1) // cols)), "white")
    for i, im in enumerate(ims):
        sheet.paste(im, ((i % cols) * w, (i // cols) * h))
    sheet.save(out_path)
    return out_path


# ---- embedded house-style definitions (generated from the original documents) ----
DOC_DEFAULTS = '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Montserrat" w:hAnsi="Montserrat"/><w:lang w:val="en-US" w:bidi="ar-SA"/></w:rPr></w:rPrDefault><w:pPrDefault/></w:docDefaults>'

STYLES = [
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:before="120" w:after="120" w:line="276" w:lineRule="auto"/><w:ind w:left="360"/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:pageBreakBefore/><w:numPr><w:numId w:val="1"/></w:numPr><w:suppressAutoHyphens/><w:spacing w:before="480" w:after="360" w:line="276" w:lineRule="auto"/><w:ind w:left="360" w:hanging="810"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat SemiBold" w:hAnsi="Montserrat SemiBold"/><w:bCs/><w:color w:val="131313"/><w:sz w:val="44"/><w:szCs w:val="56"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:numPr><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr><w:spacing w:before="360" w:after="120" w:line="276" w:lineRule="auto"/><w:ind w:left="360" w:hanging="810"/><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat Medium" w:hAnsi="Montserrat Medium"/><w:color w:val="7855FA"/><w:sz w:val="34"/><w:szCs w:val="32"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:numPr><w:ilvl w:val="2"/><w:numId w:val="1"/></w:numPr><w:spacing w:before="360" w:after="240"/><w:ind w:left="360" w:hanging="810"/><w:outlineLvl w:val="2"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat SemiBold" w:hAnsi="Montserrat SemiBold"/><w:color w:val="4E4E4E"/><w:sz w:val="26"/><w:szCs w:val="26"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:next w:val="Normal"/><w:uiPriority w:val="10"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:suppressAutoHyphens/><w:spacing w:before="3000" w:after="120" w:line="276" w:lineRule="auto"/></w:pPr><w:rPr><w:b/><w:color w:val="FFFFFF"/><w:sz w:val="60"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Subtitle"><w:name w:val="Subtitle"/><w:next w:val="Normal"/><w:uiPriority w:val="11"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:suppressAutoHyphens/><w:spacing w:after="120"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat Medium" w:hAnsi="Montserrat Medium"/><w:color w:val="FFFFFF"/><w:sz w:val="34"/><w:szCs w:val="28"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="TOCHeading"><w:name w:val="TOC Heading"/><w:basedOn w:val="Heading1"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:qFormat/><w:pPr><w:pageBreakBefore w:val="0"/><w:numPr><w:numId w:val="0"/></w:numPr><w:suppressAutoHyphens w:val="0"/><w:spacing w:before="240" w:after="0" w:line="259" w:lineRule="auto"/><w:outlineLvl w:val="9"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat SemiBold" w:hAnsi="Montserrat SemiBold"/><w:bCs w:val="0"/><w:color w:val="3907F3"/><w:sz w:val="36"/><w:szCs w:val="32"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="TOC1"><w:name w:val="toc 1"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:qFormat/><w:pPr><w:keepLines/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="11995"/></w:tabs><w:suppressAutoHyphens/><w:spacing w:before="60" w:line="276" w:lineRule="auto"/></w:pPr><w:rPr><w:noProof/><w:color w:val="7855FA"/><w:sz w:val="24"/><w:szCs w:val="28"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="TOC2"><w:name w:val="toc 2"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:qFormat/><w:pPr><w:keepLines/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="11995"/></w:tabs><w:ind w:left="360"/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="TOC3"><w:name w:val="toc 3"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:qFormat/><w:pPr><w:keepLines/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="11995"/></w:tabs><w:spacing w:line="240" w:lineRule="auto"/><w:ind w:left="720"/><w:contextualSpacing/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="Note"><w:name w:val="Note"/><w:qFormat/><w:pPr><w:keepLines/><w:shd w:val="clear" w:color="auto" w:fill="E6E6E6"/><w:spacing w:before="160" w:after="160"/><w:ind w:left="360" w:right="720"/></w:pPr><w:rPr><w:bCs/><w:szCs w:val="18"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="Instructions"><w:name w:val="Instructions"/><w:qFormat/><w:pPr><w:shd w:val="clear" w:color="auto" w:fill="FFFF00"/><w:spacing w:before="120" w:after="120"/><w:ind w:left="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat Medium" w:hAnsi="Montserrat Medium"/><w:sz w:val="26"/><w:szCs w:val="24"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="CodeBlock"><w:name w:val="Code: Block"/><w:basedOn w:val="Normal"/><w:qFormat/><w:pPr><w:pBdr><w:top w:val="single" w:sz="8" w:space="1" w:color="auto"/><w:left w:val="single" w:sz="8" w:space="4" w:color="auto"/><w:bottom w:val="single" w:sz="8" w:space="1" w:color="auto"/><w:right w:val="single" w:sz="8" w:space="4" w:color="auto"/></w:pBdr><w:shd w:val="clear" w:color="auto" w:fill="FFFFFF"/><w:spacing w:after="240" w:line="240" w:lineRule="auto"/><w:contextualSpacing/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/><w:noProof/><w:sz w:val="18"/><w:szCs w:val="22"/></w:rPr></w:style>',
    '<w:style w:type="character" w:customStyle="1" w:styleId="CodeInline"><w:name w:val="Code: Inline"/><w:basedOn w:val="DefaultParagraphFont"/><w:uiPriority w:val="1"/><w:qFormat/><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Caption"><w:name w:val="caption"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="35"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:spacing w:line="240" w:lineRule="auto"/><w:jc w:val="center"/></w:pPr><w:rPr><w:rFonts/><w:i/><w:iCs/><w:szCs w:val="18"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="Figure"><w:name w:val="Figure"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:keepLines/><w:spacing w:before="240" w:line="240" w:lineRule="auto"/><w:jc w:val="center"/></w:pPr><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="TableText"><w:name w:val="Table: Text"/><w:qFormat/><w:pPr><w:keepLines/><w:spacing w:before="40" w:after="40"/></w:pPr><w:rPr><w:szCs w:val="16"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="TableTextSmall"><w:name w:val="Table: Text Small"/><w:qFormat/><w:pPr><w:spacing w:before="80" w:after="80" w:line="276" w:lineRule="auto"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat" w:hAnsi="Montserrat"/><w:sz w:val="16"/><w:szCs w:val="17"/></w:rPr></w:style>',
    '<w:style w:type="table" w:customStyle="1" w:styleId="NutanixTable"><w:name w:val="Nutanix Table"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="99"/><w:pPr><w:spacing w:after="120"/><w:contextualSpacing/></w:pPr><w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="4E4E4E"/><w:left w:val="single" w:sz="4" w:space="0" w:color="4E4E4E"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="4E4E4E"/><w:right w:val="single" w:sz="4" w:space="0" w:color="4E4E4E"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="4E4E4E"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="4E4E4E"/></w:tblBorders><w:tblCellMar><w:top w:w="58" w:type="dxa"/><w:left w:w="58" w:type="dxa"/><w:bottom w:w="58" w:type="dxa"/><w:right w:w="58" w:type="dxa"/></w:tblCellMar></w:tblPr><w:trPr><w:cantSplit/></w:trPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="auto"/><w:vAlign w:val="center"/></w:tcPr><w:tblStylePr w:type="firstRow"><w:rPr><w:rFonts w:ascii="Montserrat SemiBold" w:hAnsi="Montserrat SemiBold"/><w:color w:val="FFFFFF"/><w:sz w:val="20"/></w:rPr><w:tblPr/><w:trPr><w:cantSplit w:val="0"/><w:tblHeader/></w:trPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="4B00AA"/></w:tcPr></w:tblStylePr></w:style>',
    '<w:style w:type="table" w:customStyle="1" w:styleId="NutanixTableSmall"><w:name w:val="Nutanix Table: Small"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="99"/><w:rPr><w:sz w:val="16"/></w:rPr><w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="888888"/><w:left w:val="single" w:sz="4" w:space="0" w:color="888888"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="888888"/><w:right w:val="single" w:sz="4" w:space="0" w:color="888888"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="888888"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="888888"/></w:tblBorders></w:tblPr><w:tcPr><w:vAlign w:val="center"/></w:tcPr><w:tblStylePr w:type="firstRow"><w:rPr><w:rFonts w:ascii="Montserrat SemiBold" w:hAnsi="Montserrat SemiBold"/><w:color w:val="FFFFFF"/></w:rPr><w:tblPr/><w:trPr><w:tblHeader/></w:trPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="4B00AA"/></w:tcPr></w:tblStylePr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListBullet"><w:name w:val="List Bullet"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:numPr><w:numId w:val="2"/></w:numPr><w:spacing w:before="120" w:after="120"/><w:ind w:left="720"/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListBullet2"><w:name w:val="List Bullet 2"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="3"/></w:numPr><w:spacing w:before="120" w:after="120"/><w:ind w:left="1080"/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListBullet3"><w:name w:val="List Bullet 3"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:numPr><w:ilvl w:val="2"/><w:numId w:val="4"/></w:numPr><w:spacing w:before="120" w:after="120"/><w:ind w:left="1440"/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListBullet4"><w:name w:val="List Bullet 4"/><w:uiPriority w:val="99"/><w:unhideWhenUsed/><w:qFormat/><w:pPr><w:numPr><w:ilvl w:val="3"/><w:numId w:val="5"/></w:numPr><w:spacing w:before="60"/><w:ind w:left="1800"/><w:contextualSpacing/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListNumber"><w:name w:val="List Number"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:numPr><w:numId w:val="6"/></w:numPr><w:spacing w:before="120" w:after="60"/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListNumber2"><w:name w:val="List Number 2"/><w:basedOn w:val="ListNumber"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:numPr><w:numId w:val="7"/></w:numPr><w:spacing w:before="60"/><w:contextualSpacing/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="ListNumber3"><w:name w:val="List Number 3"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:numPr><w:ilvl w:val="2"/><w:numId w:val="8"/></w:numPr><w:spacing w:before="60" w:after="60"/><w:ind w:left="1440"/><w:contextualSpacing/></w:pPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Header"><w:name w:val="header"/><w:uiPriority w:val="99"/><w:unhideWhenUsed/><w:qFormat/><w:pPr><w:tabs><w:tab w:val="center" w:pos="4680"/><w:tab w:val="right" w:pos="9360"/></w:tabs><w:spacing w:after="120"/></w:pPr><w:rPr><w:sz w:val="18"/></w:rPr></w:style>',
    '<w:style w:type="paragraph" w:styleId="Footer"><w:name w:val="footer"/><w:uiPriority w:val="99"/><w:qFormat/><w:pPr><w:keepLines/><w:widowControl w:val="0"/><w:tabs><w:tab w:val="right" w:pos="10800"/></w:tabs><w:autoSpaceDE w:val="0"/><w:autoSpaceDN w:val="0"/><w:adjustRightInd w:val="0"/><w:spacing w:before="120" w:after="120"/><w:ind w:left="1800"/><w:contextualSpacing/><w:jc w:val="both"/><w:textAlignment w:val="center"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat Light" w:hAnsi="Montserrat Light"/><w:color w:val="131313"/><w:sz w:val="14"/><w:szCs w:val="14"/></w:rPr></w:style>',
    '<w:style w:type="character" w:styleId="Hyperlink"><w:name w:val="Hyperlink"/><w:basedOn w:val="DefaultParagraphFont"/><w:uiPriority w:val="99"/><w:qFormat/><w:rPr><w:color w:val="391699"/><w:u w:val="single"/></w:rPr></w:style>',
]

ABSTRACT_NUMS = [
    '<w:abstractNum w:abstractNumId="1"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading1"/><w:lvlText w:val="%1"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2142" w:hanging="792"/></w:pPr><w:rPr><w:rFonts w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading2"/><w:lvlText w:val="%1.%2"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="432" w:hanging="792"/></w:pPr><w:rPr><w:rFonts w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading3"/><w:lvlText w:val="%1.%2.%3"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="432" w:hanging="792"/></w:pPr><w:rPr><w:rFonts w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="3"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3.%4"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="432" w:hanging="792"/></w:pPr><w:rPr><w:rFonts w:hint="default"/></w:rPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="2"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="38D83FC8"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:pStyle w:val="ListBullet"/><w:lvlText w:val="\uf0b7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1800" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="1" w:tplc="04090003" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2520" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="2" w:tplc="04090005" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0a7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="3240" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="3" w:tplc="04090001" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0b7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="3960" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="3"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="3454E2D6"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2160" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:hint="default"/><w:color w:val="auto"/></w:rPr></w:lvl><w:lvl w:ilvl="1" w:tplc="6166065E"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:pStyle w:val="ListBullet2"/><w:lvlText w:val="o"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2880" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="2" w:tplc="04090005" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0a7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="3600" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="3" w:tplc="04090001" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0b7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="4320" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="4"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="8286E442"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0a7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2880" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="1" w:tplc="04090003" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="3600" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="2" w:tplc="509A7702"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:pStyle w:val="ListBullet3"/><w:lvlText w:val="\uf0a7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="4320" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="3" w:tplc="04090001" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0b7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="5040" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="5"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="B720CAE4"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="–"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1080" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Montserrat" w:hAnsi="Montserrat" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="1" w:tplc="04090003" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1800" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="2" w:tplc="04090005" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\uf0a7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2520" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/></w:rPr></w:lvl><w:lvl w:ilvl="3" w:tplc="6A969B38"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:pStyle w:val="ListBullet4"/><w:lvlText w:val="\uf0b7"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="3240" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="6"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="A41E7E70"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="ListNumber"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl><w:lvl w:ilvl="1" w:tplc="04090019" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl><w:lvl w:ilvl="2" w:tplc="0409001B" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:lvlText w:val="%3."/><w:lvlJc w:val="right"/><w:pPr><w:ind w:left="2160" w:hanging="180"/></w:pPr></w:lvl><w:lvl w:ilvl="3" w:tplc="0409000F" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%4."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2880" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="7"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="66F683E0"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:pStyle w:val="ListNumber2"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl><w:lvl w:ilvl="1" w:tplc="04090019" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="2160" w:hanging="360"/></w:pPr></w:lvl><w:lvl w:ilvl="2" w:tplc="0409001B" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:lvlText w:val="%3."/><w:lvlJc w:val="right"/><w:pPr><w:ind w:left="2880" w:hanging="180"/></w:pPr></w:lvl><w:lvl w:ilvl="3" w:tplc="0409000F" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%4."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="3600" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>',
    '<w:abstractNum w:abstractNumId="8"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0" w:tplc="0FAA35EA"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:lvlText w:val="%1."/><w:lvlJc w:val="right"/><w:pPr><w:ind w:left="3870" w:hanging="360"/></w:pPr></w:lvl><w:lvl w:ilvl="1" w:tplc="04090019" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="4590" w:hanging="360"/></w:pPr></w:lvl><w:lvl w:ilvl="2" w:tplc="E026B164"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:pStyle w:val="ListNumber3"/><w:lvlText w:val="%3."/><w:lvlJc w:val="right"/><w:pPr><w:ind w:left="5310" w:hanging="180"/></w:pPr></w:lvl><w:lvl w:ilvl="3" w:tplc="0409000F" w:tentative="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%4."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="6030" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>',
]

NUMS = ['<w:num w:numId="1"><w:abstractNumId w:val="1"/></w:num>', '<w:num w:numId="2"><w:abstractNumId w:val="2"/></w:num>', '<w:num w:numId="3"><w:abstractNumId w:val="3"/></w:num>', '<w:num w:numId="4"><w:abstractNumId w:val="4"/></w:num>', '<w:num w:numId="5"><w:abstractNumId w:val="5"/></w:num>', '<w:num w:numId="6"><w:abstractNumId w:val="6"/></w:num>', '<w:num w:numId="7"><w:abstractNumId w:val="7"/></w:num>', '<w:num w:numId="8"><w:abstractNumId w:val="8"/></w:num>']

HEADER_LINE_PNG = (
    'iVBORw0KGgoAAAANSUhEUgAABGoAAAApCAIAAADiVliAAAAAAXNSR0IArs4c6QAAAAlwSFlzAAAOxAAADsQBlSsOGwAAB8lJ'
    'REFUeF7t3d1vU3UYwPG2Y32hrO3avZSVbTIGg5l1HWJiJF5o4oXemHjBnWZkY3MiL24ikeiVMbwEmBMF2QsIXPsvoDcbb8ra'
    'Duk2R7JlKGFZC3thyNisvwMETUigHvpyzq/f3dJzfs/v85zEPD7n/B5jPB438IcAAggggAACCCCAAAIIIPAsAdOzfsC/I4AA'
    'AggggAACCCCAAAIIKAKUTzwHCCCAAAIIIIAAAggggEBCApRPCTHxIwQQQAABBBBAAAEEEECA8olnAAEEEEAAAQQQQAABBBBI'
    'SIDyKSEmfoQAAggggAACCCCAAAIIUD7xDCCAAAIIIIAAAggggAACCQlQPiXExI8QQAABBBBAAAEEEEAAAaOY+3Rn6v7YlSks'
    'EEAAAQQQUCfg8dmKypYaTUZ1l3MVAggggAACehFQyqfwTxNfvtOrl4iJEwEEEEBAUwJi+vqr7/oaDtY6Cy2aCuxxMPG/438M'
    'z549PTo3s6DNCIkKAQQQQEDLAi+/7Q286c1ZovxfQqV8MsQN0Rt3tRwxsSGAAAIIaFZgNjp/as9A8Ur7pj3r8r1WDcYZ6Z3s'
    'ag3N/7VY4LNpMDxCQgABBBDQuMDr75W/tqnUlPO4fNJ4vISHAAIIIKBhAdHbCZ2dOPlpuKLO9f5XNZqqoERsw5dudbeF7M7c'
    'po5ASeUyDUMSGgIIIICADgQ4OkIHSSJEBBBAQMsC4pOn2jeK6vf7R8NTZ764MnldQ68zRM5Fj2297Cw0N31N7aTlh4jYEEAA'
    'Ad0IPHh5jz8EEEAAAQSeT0D8x2Tg54kTu8JlLzrq99a4SzL8mpyIZ+hCtEf0nVy5W9oDvjV5z7c/rkYAAQQQQEARoHziOUAA'
    'AQQQSI6AqFjEV0biTTnvSrtoRhWVL03OfVXdJdIX/X5bf2GZTZxpsXwV7+ypQuQiBBBAAIEnBHh5j4cCAQQQQCA5AkajoXpj'
    'QcNB/83RuTOfX5kYmxNHE2XkT1RxPZ+ExHHq9fv81E4ZSQGLIoAAArIK0H2SNbPsCwEEEMiYwO+XYse3B8UZEg2Har0V9nTG'
    'obyzdz56fEfQvdy6pV2sTt8pnfyshQACCMgvQPdJ/hyzQwQQQCDNApUb3KIHdXvi3uk9AzeuzabzG9tI32TnzqDHZ918wE/t'
    'lOa8sxwCCCCQDQJ0n7Ihy+wRAQQQyIDAyC+3xNdHeR6zOLmhZHXKu0DijPKhi7Hu1pCjQFmRd/YykHKWRAABBLJAgO5TFiSZ'
    'LSKAAAKZEKjckC9en7tz+/7J3eHrQzOp7kGJM8qPb+t3FVs4KyIT2WZNBBBAIFsE6D5lS6bZJwIIIJARgZFfb3XuCFrzlojJ'
    'SyuqUnJ6+KPZuK3BZfnmLcx3ykiaWRQBBBDIGgG6T1mTajaKAAIIZEKg8qV8UdLMzy10fxwc+206FSE8mo1bZKF2SgUv90QA'
    'AQQQ+K8A3SeeBwQQQACBlAtcu6z0oHKtpqaOurJqR7LW+3c2rjNX3DkNX1glK3LugwACCCCgUwG6TzpNHGEjgAACehJYtT7/'
    'g2/rFhfiR1suj4ankhW6OGfv2FbxvZO1+Qi1U7JQuQ8CCCCAwNMEKJ94PhBAAAEE0iGwstbVeKjWlGPobguJCur5T5K42jt5'
    'YlfYs8JWv7eGc/bSkULWQAABBBAwGHh5j6cAAQQQQCB9AuLzJ3E+3v17ix8eXS8KKnULi9Jr8Hy0c3u/x2drPMxsXHWKXIUA'
    'AgggoEaA7pMaNa5BAAEEEFAnUF7tEAWP1b5E9KDEB1Hi0DwV94n0TnaJ2bii77Sf2bgq/LgEAQQQQEC9AN0n9XZciQACCCCg'
    'TmD86vSxbf3zdxdbvlu/qu5/9KCU2bgXYl2tIWehWZyE7q1I+TRedRvkKgQQQAABWQXoPsmaWfaFAAIIaFegtNrR3BGwO3O7'
    'W4PDl2KJ96CU2bjb+91eq5iNS+2k3QQTGQIIICCvAN0neXPLzhBAAAFtC4xHpjt3Bu9OLzQfCaze4H56sMps3Isx8crfMre5'
    'qT2wvJK+k7azS3QIIICApAJ0nyRNLNtCAAEENC9Quk70oOryPOaunaHBc9Gnn8WnzMb9qN8pZuMepnbSfGoJEAEEEJBXgO6T'
    'vLllZwgggIAeBMYHp3vawrOx+cb22rWveJ4MWfne6WKspy1kdz2YjUvfSQ9pJUYEEEBAVgG6T7Jmln0hgAAC+hAoXeto6gg4'
    'iy3iQHMxyunJoJW+08PZuN9QO+kjp0SJAAIISCxA90ni5LI1BBBAQDcC14dmftgdvj1xb/MBf/XGAqPxUeQPZ+M6CiwNB/2+'
    'NXm62Q+BIoAAAghIKkD3SdLEsi0EEEBAVwIrqvIaDwdcRUoPKtKn9KDEp1CRvmjnjqCr2CLOKKd20lU+CRYBBBCQVoDuk7Sp'
    'ZWMIIICA7gT+HJk99dnA5Phc/T6/0WQQfSd3iW3zvhpfFX0n3SWTgBFAAAE5BSif5Mwru0IAAQR0KnBz9I7oOI0PzpgtpqIX'
    '7A9m49p1uhfCRgABBBCQT4DySb6csiMEEEBA3wKigvpx/5DZlvNWS4VvNX0nfWeT6BFAAAHJBCifJEso20EAAQQQQAABBBBA'
    'AIFUCXB0RKpkuS8CCCCAAAIIIIAAAghIJkD5JFlC2Q4CCCCAAAIIIIAAAgikSuAfHlD27kn+OSMAAAAASUVORK5CYII='
)
